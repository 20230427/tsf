#!/usr/bin/env python
"""Build/update the cumulative-ablation xlsx (openpyxl, Pamba-style pipeline).

Reads the aggregated record produced by ``scripts/ablation/collect_cumulative.py``
(``output/cumulative_ablation/cumulative_ablation.json`` -- the sole input;
per-cell provenance JSONs are never parsed here) and writes
``output/cumulative_ablation/cumulative_ablation.xlsx``.

Layout (single sheet "Cumulative Ablation", follows Pamba's
output/ablation/update_ablation.py right-block style):
  - columns: Variant | Len | one MSE/MAE pair per dataset
  - dataset column headers use the display names of documents/datasets.md
    (Weather, ECL, Exchange, Illness, Solar, ETTh*, PEMS*)
  - dataset column order: raw channel-correlation ``mean|r|`` ascending
    (``DS_PCC`` below, values from documents/datasets.md sec. 6 -- the
    dataset where channels are least correlated sits leftmost). Ties break
    by dataset name (ETTh2 before ETTm2, matching the datasets.md table);
    datasets missing from ``DS_PCC`` sort last by name. This replaces the
    earlier Full-avg-MSE ordering.
  - rows: 7 chain variants x (4 pred_lens + Avg), then a Min summary block

Update pipeline (order matters, mirrors Pamba):
  validate layout (a changed dataset set auto-rebuilds the template -- the
  xlsx is 100% derived from the verified JSON, so nothing is lost; missing
  or empty JSON stays a hard failure, fail-closed)
    -> fill cell values (seed means)
    -> fill structural no-op blocks: copy the previous effective variant
       (cumulative-chain semantics: a no-op variant is bit-identical to it)
    -> Avg rows (only when all 4 horizons are numeric)
    -> Min summary rows (no-op blocks excluded: duplicates by construction)
    -> best-cell highlight: per-column global min across non-noop variants
       -> red bold font (fonts reset first)
    -> paint no-op blocks yellow #FFFF99 last (fills must survive resets)

Idempotent: re-running after more cells land only rewrites values/styles.
Requires openpyxl (optional dependency; the markdown summary from the
collector does not need it).

Usage:
    python scripts/ablation/update_cumulative_table.py
    python scripts/ablation/update_cumulative_table.py --create
"""
from __future__ import annotations

import argparse
import json
import math
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from scripts.ablation.cumulative_registry import (  # noqa: E402
    CUMULATIVE_STEPS,
    STANDARD_PRED_LENS,
    PEMS_PRED_LENS,
    ILLNESS_PRED_LENS,
    VARIANT_DISPLAYS,
    VARIANT_NAMES,
    config_path_for,
    noop_variants_for_dataset,
    pred_lens_for,
)

HEADER_ROWS = 2
N_PL = 4
ROWS_PER_VARIANT = N_PL + 1  # 4 pred_lens + Avg
VARIANT_COL = 1
LEN_COL = 2
DATA_COL_BASE = 3
DATA_START = HEADER_ROWS + 1
N_VARIANTS = len(CUMULATIVE_STEPS)
SUMMARY_LABEL_ROW = DATA_START + N_VARIANTS * ROWS_PER_VARIANT + 1
SUMMARY_START = SUMMARY_LABEL_ROW + 1

FONT_NAME = "Aptos"
FONT_SIZE = 12
ROW_HEIGHT = 20
NUM_FMT = r"0.000_);[Red]\(0.000\)"
HEADER_FILL = "FFFFC000"
NOOP_FILL = "FFFFFF99"
BEST_COLOR = "FFCC0000"
SECTION_FONT_COLOR = "FF1F4E79"
THIN = {"style": "thin"}

DEFAULT_JSON = "output/cumulative_ablation/cumulative_ablation.json"
DEFAULT_XLSX = "output/cumulative_ablation/cumulative_ablation.xlsx"

# Display names follow documents/datasets.md (global quick-reference table):
# internal config/checkpoint stems map to the paper-style dataset names.
# Unknown stems fall back to themselves.
DATASET_DISPLAY = {
    "weather": "Weather",
    "electricity": "ECL",
    "exchange_rate": "Exchange",
    "illness": "Illness",
    "solar": "Solar",
    "traffic": "Traffic",
}

# Raw channel-correlation mean|r| per dataset -- measured values from
# documents/datasets.md sec. 6 (src/utils/channel_correlation.py; the
# aggregated output/channel_correlation/summary.csv is not tracked, so the
# table's values are vendored here the same way Pamba vendors DS_PCC).
# Keys are internal stems. Column order = ascending value.
DS_PCC = {
    "ETTh1": 0.222,
    "ETTm1": 0.224,
    "ETTh2": 0.325,
    "ETTm2": 0.325,
    "weather": 0.339,
    "electricity": 0.489,
    "exchange_rate": 0.513,
    "traffic": 0.564,
    "illness": 0.708,
    "PEMS04": 0.770,
    "PEMS08": 0.784,
    "PEMS07": 0.800,
    "PEMS03": 0.840,
    "solar": 0.916,
}


def dataset_display(dataset: str) -> str:
    return DATASET_DISPLAY.get(dataset, dataset)


def _pcc_sort_key(dataset: str):
    """Ascending raw mean|r|; ties by name (datasets.md table order);
    unknown stems sort last, then by name."""
    pcc = DS_PCC.get(dataset)
    return (0, pcc, dataset) if pcc is not None else (1, 0.0, dataset)


def _fmt3(v):
    return math.floor(v * 1000 + 0.5) / 1000


def _is_numeric(v):
    return v is not None and isinstance(v, (int, float)) and not isinstance(v, bool)


def _len_label(pi, datasets):
    """Composite horizon label: standard, plus PEMS/Illness when present
    ('96/12' or '96/24')."""
    if pi >= N_PL:
        return "Avg"
    std = STANDARD_PRED_LENS[pi]
    parts = [str(std)]
    if any(d.startswith("PEMS") for d in datasets):
        pems = PEMS_PRED_LENS[pi]
        if pems != std:
            parts.append(str(pems))
    if "illness" in datasets:
        ill = ILLNESS_PRED_LENS[pi]
        if str(ill) not in parts:
            parts.append(str(ill))
    return "/".join(parts)


def _load_results(json_path: Path):
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    aggregated = payload.get("results", {})
    datasets = sorted(aggregated.keys())
    if not datasets:
        raise SystemExit(f"no datasets in {json_path}; run collect_cumulative.py first")
    order = sorted(datasets, key=_pcc_sort_key)
    return aggregated, order


def _cell_value(aggregated, ds, variant, pl, metric):
    entry = aggregated.get(ds, {}).get(variant, {}).get(str(pl))
    if entry is None:
        return None
    return entry.get(f"{metric}_mean")


def _variant_row(vi):
    return DATA_START + vi * ROWS_PER_VARIANT


def _ds_col(di):
    return DATA_COL_BASE + 2 * di


def _create_sheet(wb, datasets):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ws = wb.active
    ws.title = "Cumulative Ablation"

    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    header_font = Font(name=FONT_NAME, bold=True, size=FONT_SIZE)
    cell_font = Font(name=FONT_NAME, bold=False, size=FONT_SIZE)
    border = Border(left=Side(**THIN), right=Side(**THIN),
                    top=Side(**THIN), bottom=Side(**THIN))
    center = Alignment(horizontal="center", vertical="center")
    vcenter = Alignment(vertical="center")

    for vi, step in enumerate(CUMULATIVE_STEPS):
        base = _variant_row(vi)
        ws.merge_cells(start_row=base, start_column=VARIANT_COL,
                       end_row=base + ROWS_PER_VARIANT - 1, end_column=VARIANT_COL)
        c = ws.cell(row=base, column=VARIANT_COL, value=step["display"])
        c.font = Font(name=FONT_NAME, bold=True, size=FONT_SIZE)
        c.alignment = vcenter
        for pi in range(ROWS_PER_VARIANT):
            lc = ws.cell(row=base + pi, column=LEN_COL,
                         value=_len_label(pi, datasets))
            lc.font = cell_font
            lc.alignment = center
            for di in range(len(datasets)):
                for off in range(2):
                    ws.cell(row=base + pi, column=_ds_col(di) + off).font = cell_font
        ws.row_dimensions[base].height = ROW_HEIGHT

    for di, ds in enumerate(datasets):
        col = _ds_col(di)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        c = ws.cell(row=1, column=col, value=dataset_display(ds))
        c.fill, c.font, c.alignment, c.border = header_fill, header_font, center, border
        for off, metric in enumerate(("MSE", "MAE")):
            mc = ws.cell(row=2, column=col + off, value=metric)
            mc.fill, mc.font, mc.alignment, mc.border = header_fill, header_font, center, border
    for row, label in ((1, "Variant"), (2, "Len")):
        c = ws.cell(row=row, column=VARIANT_COL if row == 1 else LEN_COL, value=label)
        c.fill, c.font, c.alignment, c.border = header_fill, header_font, center, border
        if row == 2:
            ws.cell(row=2, column=VARIANT_COL).fill = header_fill
            ws.cell(row=2, column=VARIANT_COL).border = border

    c = ws.cell(row=SUMMARY_LABEL_ROW, column=VARIANT_COL,
                value="Min (best across variants, no-op blocks excluded)")
    c.font = Font(name=FONT_NAME, bold=True, size=FONT_SIZE, color=SECTION_FONT_COLOR)
    for pi in range(ROWS_PER_VARIANT):
        lc = ws.cell(row=SUMMARY_START + pi, column=LEN_COL,
                     value=_len_label(pi, datasets))
        lc.font = cell_font
        lc.alignment = center
        for di in range(len(datasets)):
            for off in range(2):
                mc = ws.cell(row=SUMMARY_START + pi, column=_ds_col(di) + off)
                mc.font = cell_font
                mc.border = border

    ws.column_dimensions[get_column_letter(VARIANT_COL)].width = 24
    ws.column_dimensions[get_column_letter(LEN_COL)].width = 10
    for di in range(len(datasets)):
        ws.column_dimensions[get_column_letter(_ds_col(di))].width = 10
        ws.column_dimensions[get_column_letter(_ds_col(di) + 1)].width = 10

    # Borders on the data grid (values come later via update)
    for vi in range(N_VARIANTS):
        for pi in range(ROWS_PER_VARIANT):
            for di in range(len(datasets)):
                for off in range(2):
                    ws.cell(row=_variant_row(vi) + pi,
                            column=_ds_col(di) + off).border = border
    return ws


def _validate_layout(ws, datasets) -> str | None:
    """Return a mismatch reason (or None when the layout matches the
    current dataset set / variant chain)."""
    for di, ds in enumerate(datasets):
        if ws.cell(row=1, column=_ds_col(di)).value != dataset_display(ds):
            return (f"xlsx layout mismatch at column {_ds_col(di)}: "
                    f"expected {dataset_display(ds)!r}, found "
                    f"{ws.cell(row=1, column=_ds_col(di)).value!r}. "
                    "Dataset set changed.")
    if ws.cell(row=DATA_START, column=VARIANT_COL).value != CUMULATIVE_STEPS[0]["display"]:
        return "xlsx layout mismatch: variant block moved."
    return None


def _noop_map(datasets):
    return {ds: noop_variants_for_dataset(config_path_for(ds)) for ds in datasets}


def _fill_values(ws, aggregated, datasets):
    for vi, name in enumerate(VARIANT_NAMES):
        base = _variant_row(vi)
        for di, ds in enumerate(datasets):
            for pi in range(N_PL):
                pl = pred_lens_for(ds)[pi]
                for off, metric in enumerate(("mse", "mae")):
                    v = _cell_value(aggregated, ds, name, pl, metric)
                    c = ws.cell(row=base + pi, column=_ds_col(di) + off)
                    c.value = _fmt3(v) if v is not None else None
                    c.number_format = NUM_FMT


def _fill_noop(ws, aggregated, datasets, noop):
    """Copy no-op blocks from their previous effective variant (chain semantics)."""
    for vi, name in enumerate(VARIANT_NAMES):
        for di, ds in enumerate(datasets):
            if name not in noop[ds]:
                continue
            src = VARIANT_NAMES.index(noop[ds][name])
            base, sbase = _variant_row(vi), _variant_row(src)
            for pi in range(N_PL):
                for off in range(2):
                    c = ws.cell(row=base + pi, column=_ds_col(di) + off)
                    v = ws.cell(row=sbase + pi, column=_ds_col(di) + off).value
                    c.value = v
                    c.number_format = NUM_FMT


def _update_avg(ws, datasets):
    for vi in range(N_VARIANTS):
        base = _variant_row(vi)
        for di in range(len(datasets)):
            for off in range(2):
                vals = [ws.cell(row=base + pi, column=_ds_col(di) + off).value
                        for pi in range(N_PL)]
                nums = [v for v in vals if _is_numeric(v)]
                c = ws.cell(row=base + N_PL, column=_ds_col(di) + off)
                c.value = _fmt3(sum(nums) / N_PL) if len(nums) == N_PL else None
                c.number_format = NUM_FMT


def _update_summary(ws, datasets, noop):
    for pi in range(ROWS_PER_VARIANT):
        for di, ds in enumerate(datasets):
            for off in range(2):
                vals = []
                for vi, name in enumerate(VARIANT_NAMES):
                    if name in noop[ds]:
                        continue
                    v = ws.cell(row=_variant_row(vi) + pi,
                                column=_ds_col(di) + off).value
                    if _is_numeric(v):
                        vals.append(v)
                c = ws.cell(row=SUMMARY_START + pi, column=_ds_col(di) + off)
                c.value = min(vals) if vals else None
                c.number_format = NUM_FMT


def _highlight_best(ws, datasets, noop):
    from openpyxl.styles import Font, PatternFill

    best_font = Font(name=FONT_NAME, bold=True, size=FONT_SIZE, color=BEST_COLOR)
    cell_font = Font(name=FONT_NAME, bold=False, size=FONT_SIZE)
    default_fill = PatternFill(fill_type=None)

    for vi in range(N_VARIANTS):
        for pi in range(ROWS_PER_VARIANT):
            for col in range(DATA_COL_BASE, _ds_col(len(datasets) - 1) + 2):
                c = ws.cell(row=_variant_row(vi) + pi, column=col)
                c.font = cell_font
                c.fill = default_fill

    for pi in range(ROWS_PER_VARIANT):
        for di, ds in enumerate(datasets):
            for off in range(2):
                col = _ds_col(di) + off
                vals = []
                for vi, name in enumerate(VARIANT_NAMES):
                    if name in noop[ds]:
                        continue
                    v = ws.cell(row=_variant_row(vi) + pi, column=col).value
                    if _is_numeric(v):
                        vals.append((v, vi))
                if not vals:
                    continue
                gmin = min(v for v, _ in vals)
                for v, vi in vals:
                    if v == gmin:
                        ws.cell(row=_variant_row(vi) + pi, column=col).font = best_font


def _paint_noop(ws, datasets, noop):
    from openpyxl.styles import PatternFill

    fill = PatternFill(fill_type="solid", fgColor=NOOP_FILL)
    for vi, name in enumerate(VARIANT_NAMES):
        for di, ds in enumerate(datasets):
            if name not in noop[ds]:
                continue
            for pi in range(ROWS_PER_VARIANT):  # incl. the Avg row
                for off in range(2):
                    ws.cell(row=_variant_row(vi) + pi,
                            column=_ds_col(di) + off).fill = fill


def main() -> int:
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("openpyxl is not installed; the xlsx table is skipped "
              "(markdown summary remains available). pip install openpyxl")
        return 1

    cli = argparse.ArgumentParser(description=__doc__)
    cli.add_argument("--json", default=DEFAULT_JSON)
    cli.add_argument("--xlsx", default=DEFAULT_XLSX)
    cli.add_argument("--create", action="store_true",
                     help="rebuild the template instead of validating the layout")
    opts = cli.parse_args()

    json_path, xlsx_path = Path(opts.json), Path(opts.xlsx)
    if not json_path.is_file():
        raise SystemExit(f"missing {json_path}; run collect_cumulative.py first")

    aggregated, datasets = _load_results(json_path)
    noop = _noop_map(datasets)

    if opts.create or not xlsx_path.is_file():
        wb = Workbook()
        ws = _create_sheet(wb, datasets)
        print(f"created template {xlsx_path} (datasets: {' '.join(datasets)})")
    else:
        wb = load_workbook(xlsx_path)
        ws = wb[wb.sheetnames[0]]
        reason = _validate_layout(ws, datasets)
        if reason is not None:
            # Auto-rebuild instead of failing: every cell value, Avg, Min and
            # highlight is re-derived from the provenance-verified JSON on
            # each run, so a template built for an older dataset set holds no
            # manual edits worth preserving. (Missing/empty JSON above stays
            # a hard failure -- correctness is fail-closed upstream in
            # collect_cumulative.py.)
            print(f"[rebuild] {reason} Rebuilding the template.")
            wb = Workbook()
            ws = _create_sheet(wb, datasets)

    _fill_values(ws, aggregated, datasets)
    _fill_noop(ws, aggregated, datasets, noop)
    _update_avg(ws, datasets)
    _update_summary(ws, datasets, noop)
    _highlight_best(ws, datasets, noop)
    _paint_noop(ws, datasets, noop)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)

    noop_total = sum(len(v) for v in noop.values())
    print(f"updated {xlsx_path}: {len(datasets)} dataset(s), "
          f"{N_VARIANTS} variants, {noop_total} no-op block(s) painted")
    return 0


if __name__ == "__main__":
    sys.exit(main())
