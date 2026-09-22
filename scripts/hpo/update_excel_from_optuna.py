#!/usr/bin/env python3
"""Create and update the DD-Mamba experiment Excel from Optuna best-params JSONs.

Ported from ``references/pamba/scripts/hpo/update_excel_from_optuna.py`` and
adapted to this repo's export format (``output/hpo/DDMamba_*_best_params.json``
with dotted config keys, written by ``src/hpo/study_manager.export_results``).

Two operations:

    --create    Build ``assets/ddmamba_eff_results_<N>.xlsx`` (N = next free
                version) from ``assets/pamba_eff_results_1.xlsx``: rebuild the
                "Results" sheet with the DD-Mamba layout (14 datasets incl.
                PEMS03/04/07/08, DD-Mamba config columns), reusing literature
                baseline values, styles and the Archive sheet.
    (default)   Update the highest-N ``assets/ddmamba_eff_results_<N>.xlsx``
                from the best-params JSONs: MSE/MAE into the Results section
                (columns C-D) and best hyperparameters into the Config section.

Workbook layout (sheet "Results"):

    - Rows 1-3   : title / model names / metric headers
    - Rows 4-73  : Experiment Results (14 datasets x 4 pred_lens + Avg);
                   C/D = DD-Mamba, E-V = literature baselines, X-AA = SOTA
                   per-row array formulas (SMALL over MSE / MAE columns);
                   conditional formatting (ported from the template) paints
                   the row's 1st min red+bold and 2nd min blue+bold
    - Row 75     : "Config" section title
    - Rows 76-77 : Dataset banner / parameter headers
    - Rows 78-133: Config data (14 datasets x 4 pred_lens, no Avg rows)

Config columns (row 77), in sync with ``configs/hpo_search_spaces.yaml``
(dotted keys are looked up in the JSON ``best_params``; booleans pinned by
``forced_defaults`` such as ``model.use_revin`` are not tabulated):

    C enc_in (merged per dataset block, static) | D time_encoder
    | E mixer_layers | F mixer_placement | G d_model | H mamba_layers
    | I mamba_d_state | J freq_backbone | K freq_hidden
    | L time_kernel_size | M freq_sparsity | N dropout | O lr

PEMS baselines on create: PEMS04/PEMS08 come from the template's Archive
sheet (the "S-Mamba Paper" block; its two unlabelled model columns are
skipped), PEMS03/PEMS07 from ``assets/mani_eff_exp_3.xlsx`` (S-Mamba ..
TimesNet columns; that file's Dlinear/AutoFormer PEMS03/07 cells are
placeholder copies of Solar/Traffic rows and are left empty, as is TimeMixer
for every PEMS row).

Usage:
    python scripts/hpo/update_excel_from_optuna.py --create
    python scripts/hpo/update_excel_from_optuna.py [--dry-run] [--dataset NAME]
                                                   [--optuna-dir DIR] [--xlsx PATH]
                                                   [--list-coverage] [--no-backup]
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from copy import copy
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
ASSETS_DIR = PROJECT_ROOT / "assets"
OPTUNA_DIR = PROJECT_ROOT / "output" / "hpo"
TEMPLATE_XLSX = ASSETS_DIR / "pamba_eff_results_1.xlsx"
MANI_XLSX = ASSETS_DIR / "mani_eff_exp_3.xlsx"

SHEET_NAME = "Results"
METRIC_NUM_FMT = "0.000"
LR_NUM_FMT = "0.00E+00"

RESULTS_HEADER_ROW = 3
RESULTS_DATA_START = 4
RESULTS_DATA_END = 73
CONFIG_TITLE_ROW = 75
CONFIG_DATASET_ROW = 76
CONFIG_HEADER_ROW = 77
CONFIG_DATA_START = 78
CONFIG_DATA_END = 133

STANDARD_PRED_LENS = [96, 192, 336, 720]
PEMS_PRED_LENS = [12, 24, 48, 96]
ILLNESS_PRED_LENS = [24, 36, 48, 60]

DATASETS = {
    "ETTm1": STANDARD_PRED_LENS,
    "ETTm2": STANDARD_PRED_LENS,
    "ETTh1": STANDARD_PRED_LENS,
    "ETTh2": STANDARD_PRED_LENS,
    "ECL": STANDARD_PRED_LENS,
    "Exchange": STANDARD_PRED_LENS,
    "Weather": STANDARD_PRED_LENS,
    "Solar": STANDARD_PRED_LENS,
    "Traffic": STANDARD_PRED_LENS,
    "PEMS03": PEMS_PRED_LENS,
    "PEMS04": PEMS_PRED_LENS,
    "PEMS07": PEMS_PRED_LENS,
    "PEMS08": PEMS_PRED_LENS,
    "Illness": ILLNESS_PRED_LENS,
}

# study/experiment name (configs/*.yaml experiment.name) -> sheet dataset name
STUDY_ALIASES = {
    "ettm1": "ETTm1",
    "ettm2": "ETTm2",
    "etth1": "ETTh1",
    "etth2": "ETTh2",
    "electricity": "ECL",
    "exchange_rate": "Exchange",
    "weather": "Weather",
    "solar": "Solar",
    "traffic": "Traffic",
    "pems03": "PEMS03",
    "pems04": "PEMS04",
    "pems07": "PEMS07",
    "pems08": "PEMS08",
    "illness": "Illness",
}

PARAM_COL_MAP = {
    "model.time_encoder": "D",
    "model.channel_mixer_layers": "E",
    "model.mixer_placement": "F",
    "model.d_model": "G",
    "model.mamba_layers": "H",
    "model.mamba_d_state": "I",
    "model.freq_backbone": "J",
    "model.freq_hidden": "K",
    "model.time_kernel_size": "L",
    "model.freq_sparsity": "M",
    "model.dropout": "N",
    "train.lr": "O",
}

# Display labels for row 77 (compound names wrap inside the cell via "\n").
PARAM_HEADERS = {
    "C": "enc_in",
    "D": "time_\nencoder",
    "E": "mixer_\nlayers",
    "F": "mixer_\nplacement",
    "G": "d_model",
    "H": "mamba_\nlayers",
    "I": "mamba_\nd_state",
    "J": "freq_\nbackbone",
    "K": "freq_\nhidden",
    "L": "time_\nkernel",
    "M": "freq_\nsparsity",
    "N": "dropout",
    "O": "lr",
}

ENC_IN = {
    "ETTm1": 7, "ETTm2": 7, "ETTh1": 7, "ETTh2": 7, "ECL": 321,
    "Exchange": 8, "Weather": 21, "Solar": 137, "Traffic": 862,
    "PEMS03": 358, "PEMS04": 307, "PEMS07": 883, "PEMS08": 170, "Illness": 7,
}

EXPECTED_HEADERS = {
    "A1": "Experiment Results",
    "X1": "SOTA",
    "C2": "DD-Mamba",
    "A3": "Metric",
    "A75": "Config",
    "C77": "enc_in",
    "D77": "time_encoder",
    "G77": "d_model",
    "O77": "lr",
}

XLSX_VERSION_RE = re.compile(r"^ddmamba_eff_results_(\d+)\.xlsx$")
JSON_RE = re.compile(r"^DDMamba_(.+)_pl(\d+)_best_params\.json$")

# --- create-mode source maps -------------------------------------------------

# dataset -> (first data row, Avg row) in the template's Results sheet
TEMPLATE_DATASET_ROWS = {
    "ETTm1": (4, 8), "ETTm2": (9, 13), "ETTh1": (14, 18), "ETTh2": (19, 23),
    "ECL": (24, 28), "Exchange": (29, 33), "Weather": (34, 38),
    "Solar": (39, 43), "Traffic": (44, 48), "Illness": (49, 53),
}

# dataset -> (first data row, Avg row) in the template's Archive sheet
# ("S-Mamba Paper" block). MSE columns: F S-Mamba, H iTransformer,
# J (unknown model, skipped), L PatchTST, N Crossformer, P TiDE,
# R TimesNet, T Dlinear, V (unknown model, skipped), X AutoFormer.
ARCHIVE_PEMS_ROWS = {"PEMS08": (23, 27), "PEMS04": (35, 39)}
ARCHIVE_PEMS_COL_MAP = {"G": "F", "I": "H", "K": "L", "M": "N",
                        "O": "P", "Q": "R", "S": "T", "U": "X"}

# dataset -> (first data row, Avg row) in mani_eff_exp_3.xlsx Results sheet.
# MSE columns: E S-Mamba, G iTransformer, I PatchTST, K Crossformer,
# M TiDE, O TimesNet (Q Dlinear / S AutoFormer are placeholders for these
# two datasets -- see module docstring -- and are not copied).
MANI_PEMS_ROWS = {"PEMS03": (39, 43), "PEMS07": (49, 53)}
MANI_PEMS_COL_MAP = {"G": "E", "I": "G", "K": "I", "M": "K", "O": "M", "Q": "O"}


def _import_openpyxl():
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl is required for Excel operations. Install with:\n"
              '    pip install -e ".[tables]"')
        sys.exit(1)
    return openpyxl


def col_to_idx(col: str) -> int:
    result = 0
    for ch in col.upper():
        result = result * 26 + (ord(ch) - ord("A") + 1)
    return result


def next_col(col: str) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(col_to_idx(col) + 1)


def round_metric(value):
    return round(float(value), 3)


# -----------------------------------------------------------------------------
# create mode
# -----------------------------------------------------------------------------

def _harvest_std_baselines(old_ws):
    """{(dataset, row_idx): {col_letter: value}} for template columns E..V.

    Every metric is rounded to 3 decimals on harvest so the stored value,
    not just the display format, is a 3-decimal number."""
    out = {}
    for ds, (first, avg) in TEMPLATE_DATASET_ROWS.items():
        for r in list(range(first, avg)) + [avg]:
            vals = {}
            for col in ("E", "F", "G", "H", "I", "J", "K", "L",
                        "M", "N", "O", "P", "Q", "R", "S", "T", "U", "V"):
                v = old_ws[f"{col}{r}"].value
                if v is not None and not hasattr(v, "text"):
                    vals[col] = round_metric(v)
            out[(ds, r - first)] = vals
    return out


def _harvest_archive_pems(archive_ws):
    out = {}
    for ds, (first, avg) in ARCHIVE_PEMS_ROWS.items():
        for i, r in enumerate(list(range(first, avg)) + [avg]):
            vals = {}
            for tgt, src in ARCHIVE_PEMS_COL_MAP.items():
                mse = archive_ws[f"{src}{r}"].value
                mae = archive_ws[f"{next_col(src)}{r}"].value
                if mse is not None:
                    vals[tgt] = round_metric(mse)
                if mae is not None:
                    vals[next_col(tgt)] = round_metric(mae)
            out[(ds, i)] = vals
    return out


def _harvest_mani_pems(mani_ws):
    out = {}
    for ds, (first, avg) in MANI_PEMS_ROWS.items():
        for i, r in enumerate(list(range(first, avg)) + [avg]):
            vals = {}
            for tgt, src in MANI_PEMS_COL_MAP.items():
                mse = mani_ws[f"{src}{r}"].value
                mae = mani_ws[f"{next_col(src)}{r}"].value
                if mse is not None:
                    vals[tgt] = round_metric(mse)
                if mae is not None:
                    vals[next_col(tgt)] = round_metric(mae)
            out[(ds, i)] = vals
    return out


def _copy_style(donor, target):
    target._style = copy(donor._style)


def _sota_formula(row: int, col_letter: str) -> str:
    if col_letter in ("X", "Z"):
        rng = f"C{row}:U{row}"
        mod = 1
    else:
        rng = f"D{row}:V{row}"
        mod = 0
    k = 1 if col_letter in ("X", "Y") else 2
    return f"=SMALL(IF(MOD(COLUMN({rng}),2)={mod}, IF(ISNUMBER({rng}), {rng})), {k})"


def _write_results_sheet(wb, old_ws, std_baselines, archive_pems, mani_pems):
    openpyxl = _import_openpyxl()
    from openpyxl.worksheet.formula import ArrayFormula

    new_ws = wb.create_sheet("Results_ddmamba", 0)

    # --- header rows 1-3 (labels reused from the template) -------------------
    new_ws["A1"] = "Experiment Results"
    new_ws["X1"] = "SOTA"
    new_ws["A2"] = "Models"
    new_ws["C2"] = "DD-Mamba"
    for col in ("E", "G", "I", "K", "M", "O", "Q", "S", "U"):
        v = old_ws[f"{col}2"].value
        if v is not None:
            new_ws[f"{col}2"] = v
    new_ws["X2"] = "1st min"
    new_ws["Z2"] = "2nd min"
    new_ws["A3"] = "Metric"
    for col in ("C", "E", "G", "I", "K", "M", "O", "Q", "S", "U",
                "X", "Z"):
        new_ws[f"{col}3"] = "MSE"
        new_ws[f"{next_col(col)}3"] = "MAE"

    # --- experiment results rows 4-73 -----------------------------------------
    row = RESULTS_DATA_START
    dataset_blocks = {}
    for ds, lens in DATASETS.items():
        start = row
        for i, pl in enumerate(list(lens) + ["Avg"]):
            if i == 0:
                new_ws.cell(row=row, column=1, value=ds)
            new_ws.cell(row=row, column=2, value=pl)
            baseline_vals = (std_baselines.get((ds, i))
                             or archive_pems.get((ds, i))
                             or mani_pems.get((ds, i))
                             or {})
            for col_letter, v in baseline_vals.items():
                new_ws.cell(row=row, column=col_to_idx(col_letter), value=v)
            for col_letter in ("X", "Y", "Z", "AA"):
                ref = f"{col_letter}{row}"
                new_ws[ref] = ArrayFormula(ref=ref,
                                           text=_sota_formula(row, col_letter))
            row += 1
        dataset_blocks[ds] = (start, row - 1)

    # --- config section rows 75-133 -------------------------------------------
    new_ws.cell(row=CONFIG_TITLE_ROW, column=1, value="Config")
    new_ws.cell(row=CONFIG_DATASET_ROW, column=1, value="Dataset")
    new_ws.cell(row=CONFIG_DATASET_ROW, column=4, value="Model")
    new_ws.cell(row=CONFIG_HEADER_ROW, column=1, value="Parameter")
    for col_letter, header in PARAM_HEADERS.items():
        new_ws.cell(row=CONFIG_HEADER_ROW, column=col_to_idx(col_letter),
                    value=header)
    crow = CONFIG_DATA_START
    for ds, lens in DATASETS.items():
        cstart = crow
        for i, pl in enumerate(lens):
            if i == 0:
                new_ws.cell(row=crow, column=1, value=ds)
            new_ws.cell(row=crow, column=2, value=pl)
            new_ws.cell(row=crow, column=3, value=ENC_IN[ds])
            crow += 1
        new_ws.merge_cells(start_row=cstart, start_column=1,
                           end_row=crow - 1, end_column=1)
        new_ws.merge_cells(start_row=cstart, start_column=3,
                           end_row=crow - 1, end_column=3)

    # --- styles: copy from template donor cells --------------------------------
    for coord in ("A1", "X1", "A2", "C2", "X2", "Z2", "A3", "C3", "X3", "AA3"):
        _copy_style(old_ws[coord], new_ws[coord])
    for col_letter in ("E", "G", "I", "K", "M", "O", "Q", "S", "U", "V"):
        for r in (2, 3):
            _copy_style(old_ws[f"{col_letter}{r}"], new_ws[f"{col_letter}{r}"])

    avg_donor_row, data_donor_row = 8, 4
    for r in range(RESULTS_DATA_START, RESULTS_DATA_END + 1):
        donor_row = avg_donor_row if new_ws.cell(row=r, column=2).value == "Avg" \
            else data_donor_row
        for col in range(1, 28):  # A..AA
            _copy_style(old_ws.cell(row=donor_row, column=col),
                        new_ws.cell(row=r, column=col))
        # All metric cells (C..V) display exactly 3 decimals, zero-padded.
        for col in range(3, 23):
            new_ws.cell(row=r, column=col).number_format = METRIC_NUM_FMT

    _copy_style(old_ws["A55"], new_ws.cell(row=CONFIG_TITLE_ROW, column=1))
    _copy_style(old_ws["A56"], new_ws.cell(row=CONFIG_DATASET_ROW, column=1))
    _copy_style(old_ws["D56"], new_ws.cell(row=CONFIG_DATASET_ROW, column=4))
    _copy_style(old_ws["A57"], new_ws.cell(row=CONFIG_HEADER_ROW, column=1))
    for col in range(3, 16):  # C..O
        _copy_style(old_ws.cell(row=57, column=col),
                    new_ws.cell(row=CONFIG_HEADER_ROW, column=col))
    # Param labels wrap inside their cells (compound names carry "\n").
    from openpyxl.styles import Alignment
    for col in range(1, 16):
        new_ws.cell(row=CONFIG_HEADER_ROW, column=col).alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True)
    for r in range(CONFIG_DATA_START, CONFIG_DATA_END + 1):
        for col in range(1, 16):  # A..O
            _copy_style(old_ws.cell(row=58, column=col),
                        new_ws.cell(row=r, column=col))
        new_ws.cell(row=r, column=15).number_format = LR_NUM_FMT

    # --- merges ------------------------------------------------------------------
    new_ws.merge_cells("A1:V1")
    new_ws.merge_cells("X1:AA1")
    new_ws.merge_cells("A2:B2")
    for col in ("C", "E", "G", "I", "K", "M", "O", "Q", "S", "U"):
        new_ws.merge_cells(f"{col}2:{next_col(col)}2")
    new_ws.merge_cells("X2:Y2")
    new_ws.merge_cells("Z2:AA2")
    new_ws.merge_cells("A3:B3")
    for start, end in dataset_blocks.values():
        new_ws.merge_cells(start_row=start, start_column=1,
                           end_row=end, end_column=1)
    new_ws.merge_cells(start_row=CONFIG_TITLE_ROW, start_column=1,
                       end_row=CONFIG_TITLE_ROW, end_column=15)
    new_ws.merge_cells(start_row=CONFIG_DATASET_ROW, start_column=1,
                       end_row=CONFIG_DATASET_ROW, end_column=3)
    new_ws.merge_cells(start_row=CONFIG_DATASET_ROW, start_column=4,
                       end_row=CONFIG_DATASET_ROW, end_column=15)
    new_ws.merge_cells(start_row=CONFIG_HEADER_ROW, start_column=1,
                       end_row=CONFIG_HEADER_ROW, end_column=2)

    # --- geometry ------------------------------------------------------------------
    for r in range(1, 4):
        new_ws.row_dimensions[r].height = 15.6
    for r in range(RESULTS_DATA_START, RESULTS_DATA_END + 1):
        new_ws.row_dimensions[r].height = 13.8
    for r in range(CONFIG_TITLE_ROW, CONFIG_DATA_END + 1):
        new_ws.row_dimensions[r].height = 15.0
    new_ws.row_dimensions[CONFIG_HEADER_ROW].height = 36.0
    for col_letter, dim in old_ws.column_dimensions.items():
        if dim.width:
            new_ws.column_dimensions[col_letter].width = dim.width
    new_ws.column_dimensions["O"].width = 11.22

    _clone_sota_conditional_formatting(old_ws, new_ws)

    wb.remove(old_ws)
    new_ws.title = SHEET_NAME
    return dataset_blocks


def _clone_sota_conditional_formatting(old_ws, new_ws):
    """Port the template's per-row SOTA highlighting rules to the new layout.

    The template colors each metric cell whose value equals the row's 1st min
    (X/Y helper column) red+bold and the 2nd min (Z/AA) blue+bold; rules are
    expression-based with row-relative anchors, so the same formulas apply
    unchanged -- only the column ranges are rebuilt for rows 4-73."""
    from copy import deepcopy

    mse_cols = ("C", "E", "G", "I", "K", "M", "O", "Q", "S", "U")
    mae_cols = ("D", "F", "H", "J", "L", "N", "P", "R", "T", "V")
    mse_sqref = " ".join(f"{c}{RESULTS_DATA_START}:{c}{RESULTS_DATA_END}"
                         for c in mse_cols)
    mae_sqref = " ".join(f"{c}{RESULTS_DATA_START}:{c}{RESULTS_DATA_END}"
                         for c in mae_cols)
    for cf in old_ws.conditional_formatting:
        for rule in cf.rules:
            formula = rule.formula[0] if rule.formula else ""
            if "$X" in formula or "$Z" in formula:
                new_ws.conditional_formatting.add(mse_sqref, deepcopy(rule))
            elif "$Y" in formula or "$AA" in formula:
                new_ws.conditional_formatting.add(mae_sqref, deepcopy(rule))


def next_version_path() -> Path:
    candidates = []
    for f in ASSETS_DIR.glob("ddmamba_eff_results_*.xlsx"):
        m = XLSX_VERSION_RE.match(f.name)
        if m:
            candidates.append((int(m.group(1)), f))
    n = max((c[0] for c in candidates), default=0) + 1
    return ASSETS_DIR / f"ddmamba_eff_results_{n}.xlsx"


def create_workbook(xlsx_override=None):
    openpyxl = _import_openpyxl()
    if not TEMPLATE_XLSX.exists():
        print(f"ERROR: template not found: {TEMPLATE_XLSX}")
        return 1
    if not MANI_XLSX.exists():
        print(f"ERROR: PEMS03/07 baseline source not found: {MANI_XLSX}")
        return 1

    target = Path(xlsx_override) if xlsx_override else next_version_path()
    if not target.is_absolute():
        target = PROJECT_ROOT / target
    if target.exists():
        print(f"ERROR: target already exists: {target}")
        return 1

    wb = openpyxl.load_workbook(TEMPLATE_XLSX)
    old_ws = wb[SHEET_NAME]
    std = _harvest_std_baselines(old_ws)
    archive_pems = _harvest_archive_pems(wb["Archive"])

    mani_wb = openpyxl.load_workbook(MANI_XLSX)
    mani_ws = next((mani_wb[s] for s in mani_wb.sheetnames
                    if mani_wb[s]["A2"].value == "Models"), None)
    if mani_ws is None:
        print(f"ERROR: no 'Models' header sheet found in {MANI_XLSX}")
        return 1
    mani_pems = _harvest_mani_pems(mani_ws)

    blocks = _write_results_sheet(wb, old_ws, std, archive_pems, mani_pems)
    wb.save(target)

    n_pems_cells = sum(len(v) for v in archive_pems.values()) + \
        sum(len(v) for v in mani_pems.values())
    print(f"Created: {target}")
    print(f"  datasets: {len(blocks)} (results rows "
          f"{RESULTS_DATA_START}-{RESULTS_DATA_END}, config rows "
          f"{CONFIG_DATA_START}-{CONFIG_DATA_END})")
    print(f"  PEMS baseline cells written: {n_pems_cells}")
    print("  C/D (DD-Mamba) and Config params are empty -- run the update "
          "step (poe update) to fill them from Optuna JSONs.")
    return 0


# -----------------------------------------------------------------------------
# update mode
# -----------------------------------------------------------------------------

def validate_layout(ws) -> list:
    problems = []
    for coord, expected in EXPECTED_HEADERS.items():
        actual = ws[coord].value
        norm = "" if actual is None else str(actual).strip().replace("\n", "")
        if norm != expected:
            problems.append(f"{coord}: expected {expected!r}, got {actual!r}")
    return problems


def build_row_map(ws, start_row, end_row):
    """Scan columns A-B to build {(dataset, pred_len): row_number}."""
    row_map = {}
    current_dataset = None
    for row in range(start_row, end_row + 1):
        a_val = ws.cell(row=row, column=1).value
        b_val = ws.cell(row=row, column=2).value
        if a_val is not None:
            current_dataset = a_val
        if current_dataset and b_val is not None and isinstance(b_val, (int, float)):
            row_map[(current_dataset, int(b_val))] = row
    return row_map


def parse_json_filename(fname: str):
    m = JSON_RE.match(Path(fname).name)
    if not m:
        return None, None
    return m.group(1), int(m.group(2))


def resolve_dataset_name(raw_name: str) -> str:
    key = raw_name.lower()
    if key in STUDY_ALIASES:
        return STUDY_ALIASES[key]
    for sheet_name in DATASETS:
        if sheet_name.lower() == key:
            return sheet_name
    return raw_name


def find_json_files(optuna_dir: Path, dataset_filter=None) -> list:
    archive_dir = optuna_dir / "archive"
    files = []
    if not optuna_dir.exists():
        return files
    for f in sorted(optuna_dir.glob("DDMamba_*_best_params.json")):
        if archive_dir in f.parents:
            continue
        if dataset_filter:
            ds, _ = parse_json_filename(f.name)
            if ds is None or resolve_dataset_name(ds).lower() != dataset_filter.lower():
                continue
        files.append(f)
    return files


def find_excel_files(xlsx_override=None) -> list:
    if xlsx_override:
        p = Path(xlsx_override)
        if not p.is_absolute():
            p = PROJECT_ROOT / p
        return [p]
    candidates = []
    for f in ASSETS_DIR.glob("ddmamba_eff_results_*.xlsx"):
        m = XLSX_VERSION_RE.match(f.name)
        if m:
            candidates.append((int(m.group(1)), f))
    if not candidates:
        return []
    candidates.sort(key=lambda x: x[0], reverse=True)
    if len(candidates) > 1:
        top = candidates[0]
        print(
            f"WARNING: multiple versioned xlsx found; using {top[1].name} "
            f"(highest #{top[0]}). Ignored: {[c[1].name for c in candidates[1:]]}"
        )
    return [candidates[0][1]]


def fill_averages(ws, dry_run: bool = False) -> int:
    """Scan Avg rows; if all pred_len rows for a dataset have MSE+MAE, write the average."""
    filled = 0
    for row in range(RESULTS_DATA_START, RESULTS_DATA_END + 1):
        if ws.cell(row=row, column=2).value != "Avg":
            continue
        dataset = None
        for r in range(row - 1, RESULTS_DATA_START - 1, -1):
            a = ws.cell(row=r, column=1).value
            if a is not None:
                dataset = a
                break
        if not dataset:
            continue
        num_rows = len(DATASETS.get(dataset, STANDARD_PRED_LENS))
        data_start = row - num_rows

        for col in (3, 4):
            metric_label = ws.cell(row=RESULTS_HEADER_ROW, column=col).value
            if metric_label not in ("MSE", "MAE"):
                continue
            vals = []
            for dr in range(data_start, row):
                v = ws.cell(row=dr, column=col).value
                if v is not None:
                    vals.append(float(v))
            if len(vals) != num_rows:
                continue
            avg = round_metric(sum(vals) / len(vals))
            if not dry_run:
                c = ws.cell(row=row, column=col)
                c.value = avg
                c.number_format = METRIC_NUM_FMT
            filled += 1
        print(f"  Avg row {row} ({dataset}): C={ws.cell(row=row, column=3).value}")
    return filled


def write_param_cell(ws, row: int, col_letter: str, param_name: str, val, dry_run: bool):
    if not dry_run:
        cell = ws.cell(row=row, column=col_to_idx(col_letter))
        cell.value = val
        if param_name == "train.lr":
            cell.number_format = LR_NUM_FMT


def update_single_excel(excel_path: Path, json_files: list, dry_run: bool = False):
    openpyxl = _import_openpyxl()
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[SHEET_NAME]

    problems = validate_layout(ws)
    if problems:
        print("  ERROR: xlsx layout mismatch -- refusing to write:")
        for p in problems:
            print(f"    {p}")
        return 0, {}

    results_map = build_row_map(ws, RESULTS_DATA_START, RESULTS_DATA_END)
    config_map = build_row_map(ws, CONFIG_DATA_START, CONFIG_DATA_END)

    coverage = {}
    updated = 0
    for jf in json_files:
        dataset_raw, pred_len = parse_json_filename(jf.name)
        if dataset_raw is None:
            print(f"  SKIP (bad filename): {jf.name}")
            continue

        dataset = resolve_dataset_name(dataset_raw)
        coverage.setdefault(dataset, set()).add(pred_len)

        with open(jf) as f:
            data = json.load(f)

        attrs = data.get("best_trial_user_attrs") or {}
        mse = attrs.get("mse")
        if mse is None:
            mse = data.get("best_value")
        mae = attrs.get("mae")
        params = data.get("best_params", {})
        version = data.get("version", "")

        if mse is None:
            print(f"  SKIP (no MSE): {jf.name}")
            continue

        version_label = version if version else "legacy"
        mae_str = f"{round_metric(mae):.3f}" if mae is not None else "n/a"

        print(f"\n{jf.name}  [version={version_label}]")
        print(f"  Dataset={dataset}  pred_len={pred_len}  "
              f"MSE={round_metric(mse):.3f}  MAE={mae_str}")

        key = (dataset, pred_len)
        if key in results_map:
            row = results_map[key]
            old_mse = ws.cell(row=row, column=col_to_idx("C")).value
            old_mae = ws.cell(row=row, column=col_to_idx("D")).value
            if not dry_run:
                c = ws.cell(row=row, column=col_to_idx("C"))
                c.value = round_metric(mse)
                c.number_format = METRIC_NUM_FMT
                if mae is not None:
                    d = ws.cell(row=row, column=col_to_idx("D"))
                    d.value = round_metric(mae)
                    d.number_format = METRIC_NUM_FMT
            action = "OVERWRITE" if old_mse is not None else "NEW"
            print(f"  Results row {row}: [{action}] MSE {old_mse} -> "
                  f"{round_metric(mse):.3f}, MAE {old_mae} -> {mae_str}")
        else:
            print(f"  WARNING: no matching results row for ({dataset}, {pred_len})")

        if key in config_map:
            row = config_map[key]
            written = 0
            for param_name, col_letter in PARAM_COL_MAP.items():
                if param_name in params:
                    write_param_cell(ws, row, col_letter, param_name,
                                     params[param_name], dry_run)
                    written += 1
            print(f"  Config row {row}: updated {written}/{len(PARAM_COL_MAP)} params")
        else:
            print(f"  WARNING: no matching config row for ({dataset}, {pred_len})")

        updated += 1

    if not dry_run and updated > 0:
        print("\n--- Filling averages ---")
        avg_filled = fill_averages(ws, dry_run=dry_run)
        print(f"  Avg rows filled: {avg_filled}")

    if updated > 0 and not dry_run:
        wb.save(excel_path)
        print(f"\nExcel saved: {excel_path}")
        print(f"Total entries updated: {updated}")
    elif dry_run and updated > 0:
        print(f"\n[DRY RUN] Would update {updated} entries. No changes saved.")
    else:
        print("  No matching entries to update.")

    return updated, coverage


def print_coverage(coverage: dict):
    print(f"\n{'=' * 60}")
    print("Coverage matrix (dataset x pred_len) -- + has JSON, . missing")
    print(f"{'=' * 60}")
    total = 0
    have = 0
    for ds, lens in DATASETS.items():
        present = coverage.get(ds, set())
        marks = " ".join(f"pl{pl}{'+' if pl in present else '.'}" for pl in lens)
        cnt = sum(1 for pl in lens if pl in present)
        total += len(lens)
        have += cnt
        print(f"  {ds:<10} [{marks}]  {cnt}/{len(lens)}")
    print(f"  Total: {have}/{total} (dataset, pred_len) pairs have JSON")


def update_excel(dry_run=False, dataset_filter=None, optuna_dir=OPTUNA_DIR,
                 xlsx_override=None, list_coverage=False, no_backup=False):
    excel_files = find_excel_files(xlsx_override)
    if not excel_files:
        print("ERROR: No versioned Excel file matching "
              "'ddmamba_eff_results_<N>.xlsx' found in "
              f"{ASSETS_DIR}\nRun with --create first.")
        return

    json_files = find_json_files(optuna_dir, dataset_filter)

    if list_coverage:
        coverage = {}
        for jf in json_files:
            ds_raw, pl = parse_json_filename(jf.name)
            if ds_raw is None:
                continue
            coverage.setdefault(resolve_dataset_name(ds_raw), set()).add(pl)
        print_coverage(coverage)
        return

    if not json_files:
        print(f"No Optuna JSON files found in {optuna_dir} (excluding archive).")
        return

    total_updated = 0
    agg_coverage = {}
    for excel_path in excel_files:
        print(f"\n{'=' * 60}")
        print(f"Processing: {excel_path.name}")
        print(f"{'=' * 60}")

        if not dry_run and not no_backup:
            bak = excel_path.with_suffix(".xlsx.bak")
            shutil.copy(excel_path, bak)
            print(f"  Backup saved: {bak}")

        n, cov = update_single_excel(excel_path, json_files, dry_run=dry_run)
        total_updated += n
        for ds, lens in cov.items():
            agg_coverage.setdefault(ds, set()).update(lens)

    print(f"\nAll Excel files processed. Total entries updated: {total_updated}")
    print_coverage(agg_coverage)


def main():
    parser = argparse.ArgumentParser(
        description="Create/update experiment Excel from Optuna JSON results"
    )
    parser.add_argument("--create", action="store_true",
                        help="build assets/ddmamba_eff_results_<N>.xlsx from "
                             "the pamba template (next free version)")
    parser.add_argument("--dry-run", action="store_true",
                        help="Preview changes without saving")
    parser.add_argument("--dataset", default=None,
                        help="Only process JSON for this dataset name")
    parser.add_argument("--optuna-dir", default=str(OPTUNA_DIR),
                        help="Optuna JSON directory (default output/hpo)")
    parser.add_argument("--xlsx", default=None,
                        help="Single xlsx path (default: auto-pick highest-N "
                             "assets/ddmamba_eff_results_<N>.xlsx)")
    parser.add_argument("--list-coverage", action="store_true",
                        help="Only print dataset x pred_len coverage matrix")
    parser.add_argument("--no-backup", action="store_true",
                        help="Skip writing .bak before overwrite")
    args = parser.parse_args()

    if args.create:
        raise SystemExit(create_workbook(xlsx_override=args.xlsx))
    update_excel(
        dry_run=args.dry_run,
        dataset_filter=args.dataset,
        optuna_dir=Path(args.optuna_dir),
        xlsx_override=args.xlsx,
        list_coverage=args.list_coverage,
        no_backup=args.no_backup,
    )


if __name__ == "__main__":
    main()
