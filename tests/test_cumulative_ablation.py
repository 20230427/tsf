from __future__ import annotations

import json
import sys

import pytest
import yaml

from scripts.ablation import cumulative_registry as reg
from scripts.ablation import update_cumulative_table as uct

STANDARD_PRED_LENS = [96, 192, 336, 720]


def _write_config(tmp_path, name, model):
    path = tmp_path / f"{name}.yaml"
    path.write_text(yaml.dump({"model": model}), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

def test_resolve_overrides_accumulate(tmp_path):
    merged = reg.resolve_overrides(reg.variant_index("no_freq_enh"))
    assert merged == {
        "model.fusion": "sum",
        "model.freq_backbone": "none",
        "model.freq_sparsity": 0.0,
    }


def test_cumulative_flags_format_dotted_and_bools():
    flags = reg.cumulative_flags(reg.variant_index("no_revin"))
    assert flags[:4] == ["--model.fusion", "sum",
                         "--model.freq_backbone", "none"]
    assert "--model.use_revin" in flags
    assert flags[flags.index("--model.use_revin") + 1] == "false"


def test_is_noop_uses_recipe_values_and_model_defaults(tmp_path):
    cfg = _write_config(tmp_path, "toyA", {})
    assert reg.is_noop(cfg, reg.variant_index("no_freq_enh"))
    assert not reg.is_noop(cfg, reg.variant_index("no_gated_fusion"))
    assert not reg.is_noop(cfg, reg.variant_index("no_revin"))
    assert not reg.is_noop(cfg, 0)

    cfg = _write_config(tmp_path, "toyB",
                        {"channel_mixer_layers": 0, "freq_backbone": "fits"})
    assert reg.is_noop(cfg, reg.variant_index("no_mixer"))
    assert not reg.is_noop(cfg, reg.variant_index("no_freq_enh"))


def test_noop_copy_source_skips_over_consecutive_noops(tmp_path):
    cfg = _write_config(tmp_path, "toyC", {
        "channel_mixer_layers": 1,
        "freq_backbone": "fits",
        "time_linear_backbone": False,
        "use_revin": False,
    })
    noop = reg.noop_variants_for_dataset(cfg)
    assert noop == {"no_lin_backbone": "no_mixer", "no_revin": "no_mixer"}


def test_run_name_roundtrip():
    name = reg.run_name("ETTh1", "no_revin", 336, 2024)
    assert name == "ETTh1_cum_no_revin_h336_s2024"


def test_wave2_defaults_and_wave1_optin():
    assert reg.DEFAULT_DATASETS == ["ETTh2", "ETTm1", "PEMS08", "illness"]
    for ds in ["ETTh1", "ETTm2", "weather", "electricity", "solar",
               "exchange_rate", "PEMS04"]:
        assert ds in reg.OPTIONAL_DATASETS
    known = set(reg.DEFAULT_DATASETS) | set(reg.OPTIONAL_DATASETS)
    assert {"ETTh2", "ETTm1", "PEMS08", "illness"} <= known


def test_pred_lens_for_illness_and_pems():
    assert reg.pred_lens_for("illness") == [24, 36, 48, 60]
    assert reg.pred_lens_for("PEMS08") == [12, 24, 48, 96]
    assert reg.pred_lens_for("ETTh2") == [96, 192, 336, 720]


# ---------------------------------------------------------------------------
# Dataset display names + PCC ordering (documents/datasets.md sec. 6)
# ---------------------------------------------------------------------------

def test_dataset_display_uses_datasets_md_names():
    assert uct.dataset_display("weather") == "Weather"
    assert uct.dataset_display("electricity") == "ECL"
    assert uct.dataset_display("exchange_rate") == "Exchange"
    assert uct.dataset_display("illness") == "Illness"
    assert uct.dataset_display("solar") == "Solar"
    assert uct.dataset_display("ETTh1") == "ETTh1"
    assert uct.dataset_display("customDs") == "customDs"


def test_pcc_sort_ascending_with_tiebreak_and_unknown_last():
    order = sorted(
        ["solar", "ETTh1", "illness", "weather", "ETTm2", "ETTh2", "customDs"],
        key=uct._pcc_sort_key,
    )
    assert order == ["ETTh1", "ETTh2", "ETTm2", "weather",
                     "illness", "solar", "customDs"]


# ---------------------------------------------------------------------------
# xlsx table (create / update / auto-rebuild)
# ---------------------------------------------------------------------------

def _cell(mse, mae):
    return {"seeds": [2024], "n_seeds": 1, "mse_mean": mse, "mse_std": 0.0,
            "mae_mean": mae, "mae_std": 0.0}


def _dataset_results(variant_mse, pred_lens):
    out = {}
    for variant, mse in variant_mse.items():
        out[variant] = {
            str(pl): _cell(mse + 0.001 * i, mse + 0.01 + 0.001 * i)
            for i, pl in enumerate(pred_lens)
        }
    return out


def _payload(results):
    return {"experiment": {"kind": "cumulative_ablation"}, "results": results}


def _run_main(monkeypatch, json_path, xlsx_path, *extra):
    monkeypatch.setattr(
        sys, "argv",
        ["update_cumulative_table.py", "--json", str(json_path),
         "--xlsx", str(xlsx_path), *extra],
    )
    assert uct.main() == 0


@pytest.fixture
def table_env(tmp_path, monkeypatch):
    json_path = tmp_path / "cum.json"
    xlsx_path = tmp_path / "cum.xlsx"
    configs = tmp_path / "configs"
    configs.mkdir()

    def install(dataset, model):
        cfg = configs / f"{dataset}.yaml"
        cfg.write_text(yaml.dump({"model": model}), encoding="utf-8")
        monkeypatch.setattr(uct, "config_path_for", lambda ds: configs / f"{ds}.yaml")

    return json_path, xlsx_path, install, monkeypatch


def test_create_fills_values_avg_min_noop_and_highlight(table_env):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    json_path, xlsx_path, install, _ = table_env
    install("toyStd", {"channel_mixer_layers": 0, "freq_backbone": "fits"})

    mse = {name: 0.400 + 0.010 * i for i, name in enumerate(reg.VARIANT_NAMES)
           if name != "no_mixer"}
    json_path.write_text(json.dumps(_payload(
        {"toyStd": _dataset_results(mse, STANDARD_PRED_LENS)})), encoding="utf-8")
    _run_main(table_env[3], json_path, xlsx_path)

    ws = load_workbook(xlsx_path).active
    assert ws.cell(row=1, column=3).value == "toyStd"

    full_row, mixer_row, fenh_row = 3, 3 + 3 * 5, 3 + 2 * 5
    assert ws.cell(row=full_row, column=3).value == 0.4
    avg = uct._fmt3(sum(0.400 + 0.001 * i for i in range(4)) / 4)
    assert ws.cell(row=full_row + 4, column=3).value == avg

    mixer_cell = ws.cell(row=mixer_row, column=3)
    assert mixer_cell.value == ws.cell(row=fenh_row, column=3).value
    assert mixer_cell.fill.fgColor.rgb == "FFFFFF99"

    per_h = [ws.cell(row=3 + v * 5 + p, column=3).value
             for v, name in enumerate(reg.VARIANT_NAMES)
             for p in [0] if name != "no_mixer"]
    assert ws.cell(row=uct.SUMMARY_START, column=3).value == min(per_h)

    best = ws.cell(row=full_row, column=3).font
    assert best.bold and best.color.rgb == "FFCC0000"


def test_avg_blank_when_any_horizon_missing(table_env):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    json_path, xlsx_path, install, monkey = table_env
    install("toyStd", {})

    results = _dataset_results(
        {name: 0.400 + 0.010 * i for i, name in enumerate(reg.VARIANT_NAMES)},
        STANDARD_PRED_LENS,
    )
    del results["no_gated_fusion"]["336"]
    json_path.write_text(json.dumps(_payload({"toyStd": results})), encoding="utf-8")
    _run_main(monkey, json_path, xlsx_path)

    ws = load_workbook(xlsx_path).active
    assert ws.cell(row=3 + 4, column=3).value is not None
    assert ws.cell(row=3 + 1 * 5 + 4, column=3).value is None


def test_layout_change_auto_rebuilds_instead_of_failing(table_env, capsys):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    json_path, xlsx_path, install, monkey = table_env
    install("toyStd", {})
    install("toyNew", {})

    mse = {name: 0.400 + 0.010 * i for i, name in enumerate(reg.VARIANT_NAMES)}
    json_path.write_text(json.dumps(_payload(
        {"toyStd": _dataset_results(mse, STANDARD_PRED_LENS)})), encoding="utf-8")
    _run_main(monkey, json_path, xlsx_path)

    json_path.write_text(json.dumps(_payload({
        "toyStd": _dataset_results(mse, STANDARD_PRED_LENS),
        "toyNew": _dataset_results(
            {k: v - 0.05 for k, v in mse.items()}, STANDARD_PRED_LENS),
    })), encoding="utf-8")
    _run_main(monkey, json_path, xlsx_path)

    assert "[rebuild]" in capsys.readouterr().out
    ws = load_workbook(xlsx_path).active
    headers = [ws.cell(row=1, column=3 + 2 * i).value for i in range(2)]
    assert headers == ["toyNew", "toyStd"]


def test_column_order_follows_pcc_and_renames_headers(table_env):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    json_path, xlsx_path, install, monkey = table_env
    install("ETTh1", {})
    install("weather", {})
    install("PEMS04", {})

    mse = {name: 0.400 + 0.010 * i for i, name in enumerate(reg.VARIANT_NAMES)}
    json_path.write_text(json.dumps(_payload({
        "weather": _dataset_results(mse, STANDARD_PRED_LENS),
        "PEMS04": _dataset_results({k: 0.100 for k in mse}, [12, 24, 48, 96]),
        "ETTh1": _dataset_results(mse, STANDARD_PRED_LENS),
    })), encoding="utf-8")
    _run_main(monkey, json_path, xlsx_path)

    ws = load_workbook(xlsx_path).active
    headers = [ws.cell(row=1, column=3 + 2 * i).value for i in range(3)]
    assert headers == ["ETTh1", "Weather", "PEMS04"]
    assert ws.cell(row=3, column=3 + 4).value == 0.1
    assert ws.cell(row=3, column=2).value == "96/12"


def test_update_is_idempotent(table_env):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook
    json_path, xlsx_path, install, monkey = table_env
    install("toyStd", {"channel_mixer_layers": 0, "freq_backbone": "fits"})

    mse = {name: 0.4051 + 0.0107 * i for i, name in enumerate(reg.VARIANT_NAMES)
           if name != "no_mixer"}
    json_path.write_text(json.dumps(_payload(
        {"toyStd": _dataset_results(mse, STANDARD_PRED_LENS)})), encoding="utf-8")
    _run_main(monkey, json_path, xlsx_path)
    ws1 = load_workbook(xlsx_path).active
    first = [ws1.cell(row=r, column=c).value
             for r in range(1, uct.SUMMARY_START + 5) for c in (3, 4)]

    _run_main(monkey, json_path, xlsx_path)
    ws2 = load_workbook(xlsx_path).active
    second = [ws2.cell(row=r, column=c).value
              for r in range(1, uct.SUMMARY_START + 5) for c in (3, 4)]
    assert first == second


def test_empty_results_fail_closed(table_env):
    json_path, xlsx_path, install, _ = table_env
    install("toyStd", {})
    json_path.write_text(json.dumps(_payload({})), encoding="utf-8")
    with pytest.raises(SystemExit, match="no datasets"):
        uct._load_results(json_path)
